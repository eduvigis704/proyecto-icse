import streamlit as st
from docx import Document
import json
import os
from io import BytesIO

# --- CONFIGURACIÓN DE PÁGINA ---
st.set_page_config(page_title="Proyecto Intermodular ICSE", page_icon="🔬", layout="centered")

# --- ESTILOS PERSONALIZADOS (Colores ICSE) ---
st.markdown("""
    <style>
    .main { background-color: #13263d; color: white; }
    .stButton>button { 
        background-color: #0569b3; color: white; border-radius: 10px; 
        border: 2px solid #ffffff; height: 3em; width: 100%; font-weight: bold;
    }
    .stButton>button:hover { border: 2px solid #ec6216; color: #ec6216; }
    div.stTextArea textarea { border: 2px solid #0569b3; border-radius: 10px; }
    .css-10trblm { color: #ec6216 !important; } /* Títulos */
    .header-icse { background-color: white; padding: 10px; border-radius: 0 0 15px 15px; margin-bottom: 20px; text-align: center; }
    </style>
    """, unsafe_allow_html=True)

# --- BASE DE DATOS DE CONTENIDO ---
PROYECTOS = {
    "proyecto1": {
        "nombre": "Laboratorio de Diagnóstico Microbiológico",
        "mision": "Diseñar un laboratorio centrado en el cultivo, aislamiento e identificación de microorganismos.",
        "fases": [
            {
                "id": "f1", "nombre": "Fase 1: Identificación de necesidades",
                "preguntas": [
                    {
                        "id": "p1", "texto": "¿Por qué es necesario este servicio tras el aumento de infecciones nosocomiales?",
                        "pistas": ["Busca la definición de laboratorio de cultivo.", "Relaciónalo con infecciones intrahospitalarias.", "Piensa en resistencias bacterianas.", "Impacto en el SCS.", "Eficiencia en costes."]
                    }
                ]
            }
        ]
    },
    "proyecto2": {"nombre": "Laboratorio de Biología Molecular", "mision": "Diseñar un laboratorio de extracción y PCR.", "fases": []},
    "proyecto3": {"nombre": "Laboratorio de Bioquímica y Hematología", "mision": "Diseñar un laboratorio de alta rotación.", "fases": []}
}

# --- FUNCIONES DE PERSISTENCIA ---
def cargar_datos():
    if os.path.exists("db_proyectos.json"):
        with open("db_proyectos.json", "r", encoding="utf-8") as f:
            return json.load(f)
    return {"proyecto1": {}, "proyecto2": {}, "proyecto3": {}}

def guardar_datos(datos):
    with open("db_proyectos.json", "w", encoding="utf-8") as f:
        json.dump(datos, f, indent=4)

def generar_docx(usuario, datos_usuario):
    doc = Document()
    doc.add_heading(f"Proyecto Intermodular - {usuario.upper()}", 0)
    for fase in PROYECTOS[usuario]["fases"]:
        doc.add_heading(fase["nombre"], level=1)
        for preg in fase["preguntas"]:
            doc.add_heading(preg["texto"], level=2)
            res = datos_usuario.get(preg["id"], "Sin respuesta.")
            doc.add_paragraph(res)
    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()

# --- LÓGICA DE INTERFAZ ---
if 'logged_in' not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.user = ""

# Cabecera Blanca ICSE
st.markdown('<div class="header-icse"><img src="https://icse.es/wp-content/uploads/2021/04/logo-icse.png" width="160"></div>', unsafe_allow_html=True)

if not st.session_state.logged_in:
    st.title("Inicio de Sesión - CDI BioTech")
    user_input = st.text_input("Usuario (proyecto1, proyecto2, proyecto3 o admin_icse):").lower().strip()
    if st.button("ENTRAR"):
        if user_input in ["proyecto1", "proyecto2", "proyecto3", "admin_icse"]:
            st.session_state.logged_in = True
            st.session_state.user = user_input
            st.rerun()
        else:
            st.error("Usuario no válido.")
else:
    # --- PANEL DE PROFESORA (ADMIN) ---
    if st.session_state.user == "admin_icse":
        st.title("👨‍🏫 Panel de Control Docente")
        db = cargar_datos()
        for p_id, p_info in PROYECTOS.items():
            with st.expander(f"Ver progreso: {p_info['nombre']}"):
                respuestas = db.get(p_id, {})
                if not respuestas: st.write("Aún no hay respuestas.")
                for fid, res in respuestas.items():
                    st.info(f"**Pregunta {fid}:** {res}")
        if st.button("Cerrar Sesión"):
            st.session_state.logged_in = False
            st.rerun()

    # --- PANEL DE ALUMNO ---
    else:
        user = st.session_state.user
        db = cargar_datos()
        
        st.sidebar.title(f"Hola, {user}")
        if st.sidebar.button("Cerrar Sesión (Pausa)"):
            st.warning("Recuerda que para guardar en el .docx final debes pulsar 'Guardar Fase'.")
            st.session_state.logged_in = False
            st.rerun()

        st.title(f"🚀 {PROYECTOS[user]['nombre']}")
        st.markdown(f"**Misión:** {PROYECTOS[user]['mision']}")
        
        tab1, tab2 = st.tabs(["Macro-Reto", "Fases del Proyecto"])
        
        with tab1:
            st.write("El Servicio Canario de la Salud requiere la puesta en marcha del complejo...")
        
        with tab2:
            fase_sel = st.selectbox("Selecciona la Fase:", [f["nombre"] for f in PROYECTOS[user]["fases"]])
            fase_data = next(f for f in PROYECTOS[user]["fases"] if f["nombre"] == fase_sel)
            
            for p in fase_data["preguntas"]:
                st.subheader(p["texto"])
                # Recuperar respuesta guardada
                val_previo = db[user].get(p["id"], "")
                res_alumno = st.text_area("Escribe aquí tu respuesta:", value=val_previo, key=p["id"])
                
                # Sistema de Pistas
                if f"hint_step_{p['id']}" not in st.session_state: st.session_state[f"hint_step_{p['id']}"] = 0
                
                col1, col2 = st.columns([1, 2])
                with col1:
                    if st.button(f"💡 Pista (?)", key=f"btn_{p['id']}"):
                        if st.session_state[f"hint_step_{p['id']}"] < 5:
                            st.session_state[f"hint_step_{p['id']}"] += 1
                
                step = st.session_state[f"hint_step_{p['id']}"]
                if step > 0:
                    st.warning(p["pistas"][step-1])
                    if step < 5: st.info("Pulsa el botón de nuevo para la siguiente pista.")
                    else: st.error("No hay más pistas disponibles.")

            if st.button("💾 GUARDAR FASE Y ACTUALIZAR DOCUMENTO"):
                # Sincronizar y guardar
                for p in fase_data["preguntas"]:
                    db[user][p["id"]] = st.session_state[p["id"]]
                guardar_datos(db)
                st.success("Progreso guardado en el servidor.")
                
                # Botón de descarga del Word
                docx_file = generar_docx(user, db[user])
                st.download_button(label="📥 Descargar Documento .docx para entregar", 
                                   data=docx_file, 
                                   file_name=f"Proyecto_{user}.docx", 
                                   mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")